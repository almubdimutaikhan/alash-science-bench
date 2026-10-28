"""Convert a filled-in Alash_review.xlsx into review/human_reviews/<name>.csv.

    python3 scripts/import_excel_reviews.py FILLED.xlsx [--out review/human_reviews]

Needs openpyxl. Reads the reviewer's name from the instruction sheet, maps the Kazakh drop-down
labels back to the English codes used in review/review_sheet.csv, and prints a summary
(items answered, verdict counts, hours logged) plus anything it could not map.
"""
import argparse
import csv
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_pack_common import *  # noqa: E402,F403
from make_review_pack import SHEET_COLUMNS  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
CSV = ROOT / 'data' / 'question_pairs.csv'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx', type=Path)
    parser.add_argument('--out', type=Path, default=ROOT / 'review' / 'human_reviews')
    args = parser.parse_args()
    wb = load_workbook(args.xlsx, data_only=False)
    name = str(wb[SHEET_INSTR][NAME_CELL].value or '').strip()
    if not name:
        raise SystemExit('Reviewer name (sheet Нұсқаулық, cell B3) is empty.')
    with CSV.open(encoding='utf-8', newline='') as f:
        data = {r['id']: r for r in csv.DictReader(f)}
    problems, out_rows = [], {}
    q, c = wb[SHEET_Q], wb[SHEET_CHECK]
    for row in range(2, LAST_ROW + 1):
        cid = q.cell(row, 1).value
        if cid not in data:
            problems.append(f'sheet 1 row {row}: unknown ID {cid!r}')
            continue
        d = data[cid]
        rec = {k: '' for k in SHEET_COLUMNS}
        rec.update({k: d[k] for k in ('id', 'field', 'author', 'work', 'question')})
        for i, letter in enumerate('ABCD', 1):
            rec[f'choice_{letter}'] = d[f'mc_answer{i}']
        rec['reviewer'] = name
        rec['blind_short_answer'] = str(q.cell(row, 9).value or '').strip()
        rec['blind_mcq_letter'] = str(q.cell(row, 10).value or '').strip()
        rec['checked_against_scan'] = 'no'
        out_rows[cid] = (rec, row)
    for row in range(2, LAST_ROW + 1):
        cid = c.cell(row, 1).value
        if cid not in out_rows:
            problems.append(f'sheet 2 row {row}: ID {cid!r} not on sheet 1')
            continue
        rec = out_rows[cid][0]
        for col, (key, mapping) in CHECK_CHOICES.items():
            value = c[f'{col}{row}'].value
            if value in (None, ''):
                continue
            value = str(value).strip()
            if value in mapping:
                rec[key] = mapping[value]
            else:
                problems.append(f'{cid}: unrecognised {key} value {value!r}')
        rec['notes'] = str(c[f'O{row}'].value or '').strip()
    hours = 0.0
    t = wb[SHEET_TIME]
    for row in range(2, 201):
        b, e = t.cell(row, 2).value, t.cell(row, 3).value
        if hasattr(b, 'hour') and hasattr(e, 'hour'):
            hours += ((e.hour * 60 + e.minute) - (b.hour * 60 + b.minute)) / 60
    args.out.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r'[^\w-]+', '_', name, flags=re.UNICODE).strip('_').lower() or 'reviewer'
    dest = args.out / f'{slug}.csv'
    with dest.open('w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=SHEET_COLUMNS, lineterminator='\n')
        w.writeheader()
        w.writerows(rec for rec, _ in out_rows.values())
    recs = [r for r, _ in out_rows.values()]
    print(f'wrote {dest} ({len(recs)} rows) for reviewer {name!r}')
    print('blind answers:', sum(bool(r['blind_mcq_letter']) for r in recs), '/ 160')
    print('verdicts:', dict(Counter(r['verdict'] or '(none)' for r in recs)))
    print(f'hours logged (hh:mm entries only): {hours:.2f}')
    for p in problems:
        print('PROBLEM:', p)


if __name__ == '__main__':
    main()
